"""
Genera gesp_ft_14_ficticio.docx llenando la plantilla VACÍA del formato
GESP-FT-14 V3 con datos 100 % FICTICIOS (ninguna persona real).

    python tests/fixtures/generar_ficticio.py <plantilla_vacia.docx> <salida.docx>

El PDF equivalente se obtiene exportando el .docx desde Word (Guardar como PDF).
"""
import copy, re, sys, unicodedata, docx
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


def norm(s):
    s = unicodedata.normalize("NFKD", s or "")
    s = "".join(c for c in s if not unicodedata.combining(c))
    return re.sub(r"\s+", " ", s).strip().lower()


d = docx.Document(sys.argv[1])


def todas_tablas(tablas):
    for t in tablas:
        yield t
        for row in t.rows:
            for c in row.cells:
                yield from todas_tablas(c.tables)


T = lambda: list(todas_tablas(d.tables))


def celdas_unicas(row):
    out, prev = [], None
    for c in row.cells:
        if c._tc is not prev:
            out.append(c)
        prev = c._tc
    return out


def poner(celda, texto):
    p = celda.paragraphs[0]
    for r in p.runs[1:]:
        r._r.getparent().remove(r._r)
    if p.runs:
        p.runs[0].text = texto
    else:
        p.add_run(texto)


def filas_con(patron, tablas=None):
    res = []
    for t in (tablas or T()):
        for row in t.rows:
            cs = celdas_unicas(row)
            for i, c in enumerate(cs):
                if not c.tables and re.search(patron, norm(c.text)):
                    res.append((cs, i))
    return res


def derecha(patron, valor, k=0, desp=1, tablas=None):
    cs, i = filas_con(patron, tablas)[k]
    poner(cs[i + desp], valor)


def izquierda(patron, valor="X", k=0, tablas=None):
    cs, i = filas_con(patron, tablas)[k]
    poner(cs[i - 1], valor)


def texto_tras(patron, texto, k=0):
    hits = [p for p in d.element.body.iter(qn("w:p"))
            if re.search(patron, norm("".join(t.text or "" for t in p.iter(qn("w:t")))))]
    p = hits[k]
    nuevo = copy.deepcopy(p)
    ts = list(nuevo.iter(qn("w:t")))
    for t in ts:
        t.text = ""
    ts[0].text = texto
    for b in list(nuevo.iter(qn("w:b"))):
        b.getparent().remove(b)
    p.addnext(nuevo)


# I. Información general
t0 = d.tables[0]
vals = celdas_unicas(t0.rows[3])
for idx, v in zip([1, 2, 3, 5, 6, 7], ["05", "02", "2026", "20", "03", "2026"]):
    poner(vals[idx], v)
poner(celdas_unicas(t0.rows[1])[-1], "OT-2026-0457")
izquierda(r"^reevaluacion de nivel del riesgo por hechos")

# II. Datos personales
derecha(r"^cedula de ciudadania$", "X")
derecha(r"^numero de identificacion$", "1.234.567.890")
derecha(r"^primer nombre$", "PEDRO")
derecha(r"^segundo nombre$", "ANDRÉS")
derecha(r"^primer apellido$", "PÉREZ")
derecha(r"^segundo apellido$", "GÓMEZ")
derecha(r"^enfoque diferencial e interseccionalidad$", "Campesino, víctima del conflicto armado")

# Solicitante
derecha(r"^fecha de solicitud", "10/01/2026")
derecha(r"^solicitante:", "A nombre propio")
derecha(r"^tipo de evaluacion de riesgo:", "Reevaluación por hechos sobrevinientes")
texto_tras(r"^sinopsis de la informacion:$",
           "El evaluado reporta amenazas recibidas por llamada telefónica en enero de 2026.", 0)

# Antecedentes
t_ant = [t for t in T() if any(norm(c.text) == "tramite de emergencia" for c in t.rows[0].cells)][0]
for c, v in zip(celdas_unicas(t_ant.rows[1]),
                ["OT-2024-0123", "NO", "Individual", "Extraordinario", "Esquema tipo 1",
                 "Amenazas 2024", "Res. 0456 de 2024"]):
    poner(c, v)
texto_tras(r"^sinopsis de la informacion:$",
           "Cuenta con evaluación previa en 2024 con nivel extraordinario.", 1)

# Perfil antiguo (PPR)
derecha(r"^fecha de ingreso a las antiguas farc", "1998")
derecha(r"^zona de operacion$", "Sur del Tolima")
derecha(r"^columna, frente", "Frente 21")
derecha(r"^bloque donde estuvo", "Bloque Central")
derecha(r"^seudonimo$", "El Profe")
derecha(r"^rol$", "Guerrillero")
derecha(r"^actividad$", "Radista")

# Perfil actual
derecha(r"^edad actual$", "48")
derecha(r"^departamento de residencia$", "Tolima")
derecha(r"^municipio de residencia$", "Planadas")
derecha(r"^¿vive en zona rural o urbana", "Rural")
derecha(r"^¿vive en zona de reserva campesina", "Si")
derecha(r"^¿pertenece a un resguardo", "No")
derecha(r"^nivel de escolaridad$", "Bachiller")

t_fam = [t for t in T() if norm(t.rows[0].cells[0].text).startswith("grupo familiar")][0]
fam = [("MARÍA LÓPEZ", "45", "Compañera permanente", "Planadas"),
       ("JUAN PÉREZ LÓPEZ", "15", "Hijo", "Planadas"),
       ("ROSA GÓMEZ", "72", "Madre", "Chaparral")]
for r, f in zip(t_fam.rows[2:], fam):
    cs = celdas_unicas(r)
    for c, v in zip(cs[1:], f):
        poner(c, v)

derecha(r"^fuentes principal de ingresos$", "Proyecto productivo")
derecha(r"^¿se encuentra empleado", "No")
derecha(r"^¿el proyecto productivo se encuentra activo", "Si")
derecha(r"^¿participa en toar", "Si")
derecha(r"^¿cual\?$", "Desminado humanitario")
derecha(r"^¿comparece ante la jurisdiccion", "Si")
derecha(r"^maximo responsable$", "X")
derecha(r"^¿pertenece a alguna organizacion social", "Si")
derecha(r"^nombre de la colectividad$", "Partido Comunes")

t_desp = [t for t in T() if any(norm(c.text) == "motivo desplazamiento" for r in t.rows[:2] for c in r.cells)][0]
cs = celdas_unicas(t_desp.rows[2])
for c, v in zip(cs, ["Planadas, Tolima", "Ibagué, Tolima", "Terciaria, mal estado", "Moto",
                     "Diurno", "Una vez al mes", "Laboral"]):
    poner(c, v)


# Hechos de riesgo: #1 y duplicado #2
def llenar_hecho(t, dia, mes, anio, dep, mun, actor_lbl, nombre, victima, medio, tipo_lbl,
                 motiv, relato, nexo):
    ts = [t]
    for pat, v in ((r"^dia$", dia), (r"^mes$", mes), (r"^ano$", anio)):
        cs, i = filas_con(pat, ts)[0]
        poner(cs[i + 1], v)
    derecha(r"^departamento$", dep, tablas=ts)
    derecha(r"^municipio$", mun, tablas=ts)
    izquierda(actor_lbl, tablas=ts)
    derecha(r"^nombre del actor generador", nombre, tablas=ts)
    derecha(r"^victima del hecho", victima, tablas=ts)
    derecha(r"^medio hecho de riesgo$", medio, tablas=ts)
    izquierda(tipo_lbl, tablas=ts)
    derecha(r"^motivacion del hecho", motiv, tablas=ts)
    derecha(r"^relato abierto", relato, tablas=ts)
    izquierda(r"^si$" if nexo else r"^no$", tablas=ts)


t_h1 = [t for t in T() if norm(t.rows[0].cells[0].text) == "hecho de riesgo #1"][0]
t_h2_xml = copy.deepcopy(t_h1._tbl)
llenar_hecho(t_h1, "12", "01", "2026", "Tolima", "Planadas", r"^gao-r$", "Disidencias Frente 21",
             "Evaluado", "Llamada", r"^directa$", "Su rol como firmante de paz",
             "Recibió una llamada donde le advierten que debe abandonar la zona.", True)
t_h1._tbl.addnext(t_h2_xml)
t_h1._tbl.addnext(OxmlElement("w:p"))
for t in t_h2_xml.iter(qn("w:t")):
    if t.text and "#1" in t.text:
        t.text = t.text.replace("#1", "#2")
t_h2 = docx.table.Table(t_h2_xml, t_h1._parent)
llenar_hecho(t_h2, "", "11", "2025", "Tolima", "Chaparral", r"^otro$", "Desconocido",
             "Familiar", "Panfleto", r"^potencial$", "Estigmatización",
             "Apareció un panfleto en la vereda mencionando a firmantes de paz.", False)
texto_tras(r"^obsevaciones:$", "Hechos denunciados ante Fiscalía.")

# Verificaciones
for k, (lbl, nombre) in enumerate([(r"^terceros$", "Presidente JAC vereda La Esperanza"),
                                   (r"^institucion del estado colombiano$",
                                    "Personería Municipal de Planadas")]):
    t_v = [t for t in T() if norm(t.rows[0].cells[0].text) == f"verificacion #{k+1}"][0]
    derecha(lbl, "X", tablas=[t_v])
    derecha(r"^nombre de la fuente", nombre, tablas=[t_v])
    trs = [r._tr for r in t_v.rows]
    cs, i = filas_con(r"^sinopsis de la verificacion$", [t_v])[0]
    fila_sig = t_v.rows[trs.index(cs[0]._tc.getparent()) + 1]
    poner(celdas_unicas(fila_sig)[0], f"La fuente {k+1} confirma los hechos narrados por el evaluado.")


# Impacto consecuencial
def impacto(patron, si=True):
    cs, i = filas_con(patron)[-1]
    poner(cs[i + (1 if si else 2)], "X")


impacto(r"^dependencia de subsidios", True)
impacto(r"^perdida de iniciativas", False)
impacto(r"^ruptura del tejido", True)
impacto(r"^restriccion de movilidad", True)
impacto(r"^estigmatizacion$", True)
impacto(r"^afectacion psicosocial$", True)
impacto(r"^desescolarizacion$", False)


def seccion(titulo, texto, fila=1):
    t = [t for t in d.tables if norm(t.rows[0].cells[0].text) == titulo][0]
    poner(celdas_unicas(t.rows[fila])[0], texto)


seccion("vulnerabilidades y capacidades del colectivo", "Vive en zona rural dispersa con baja presencia institucional.")
seccion("contexto de orden publico", "En el sur del Tolima hay presencia de disidencias que disputan corredores de movilidad.")
seccion("alertas tempranas", "AT 012-2025 de la Defensoría del Pueblo para Planadas.")
seccion("medidas de emergencia", "Trámite de emergencia activado el 15/01/2026: medio de comunicación.")
seccion("medidas de proteccion vigentes", "Esquema tipo 1 según Resolución 0456 de 2024.")
seccion("nivel del riesgo", "EXTRAORDINARIO")
t_nr = [t for t in d.tables if norm(t.rows[0].cells[0].text) == "nivel del riesgo"][0]
poner(celdas_unicas(t_nr.rows[3])[0], "Se concluye riesgo extraordinario por amenaza directa y contexto territorial.")
seccion("modificaciones solicitadas en premesa por los y las delegadas de la mtsp", "Ninguna.")

# Analista
derecha(r"^nombres y apellidos$", "ANALISTA DE PRUEBA")
derecha(r"^documento de identidad", "99.999.999")
derecha(r"^cargo o rol$", "Analista de riesgo")
derecha(r"^correo electronico institucional$", "analista.prueba@ejemplo.gov.co")

d.save(sys.argv[2])
print("ok")
